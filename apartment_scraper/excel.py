from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Listing

# (header, attribute, width, number format)
COLUMNS = [
    ("Price €", "price", 12, "#,##0"),
    ("Rooms", "rooms", 7, "0"),
    ("Area m²", "area", 9, "0.0"),
    ("€/m²", "price_per_m2", 9, "#,##0"),
    ("District", "district", 13, None),
    ("Street", "street", 30, None),
    ("Floor", "floor", 8, None),
    ("Series", "series", 12, None),
    ("Source", "source", 10, None),
    ("Link", "url", 14, None),
]


def write_xlsx(results: list[Listing], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Apartments"
    ws.append([c[0] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")

    link_col = [c[1] for c in COLUMNS].index("url") + 1
    for row, listing in enumerate(results, start=2):
        for col, (_, attr, _, fmt) in enumerate(COLUMNS, start=1):
            value = getattr(listing, attr)
            cell = ws.cell(row=row, column=col, value=value)
            if fmt:
                cell.number_format = fmt
        if listing.url:
            cell = ws.cell(row=row, column=link_col, value="Open listing")
            cell.hyperlink = listing.url
            cell.font = Font(color="0563C1", underline="single")

    for col, (_, _, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
